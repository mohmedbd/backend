import logging

from rest_framework import viewsets
from rest_framework.permissions import AllowAny
from rest_framework_simplejwt.authentication import JWTAuthentication
from django.http import HttpResponse
from rest_framework.decorators import api_view, permission_classes, authentication_classes
from rest_framework.response import Response
from rest_framework import status
from django.utils import timezone
from django.db.models import Q
from datetime import timedelta

logger = logging.getLogger(__name__)
from rest_framework_simplejwt.tokens import AccessToken as AT

from .models import Departement, Filiere, Semestre, Module, Salle, Planification, SeanceQR, QRCode, Presence, AuditLog
from .serializers import (
    DepartementSerializer, FiliereSerializer, SemestreSerializer, ModuleSerializer, SalleSerializer,
    PlanificationSerializer, SeanceQRSerializer,
    QRCodeSerializer, PresenceSerializer,
    StudentUserSerializer, ProfessorUserSerializer,
    HeadOfDepartmentUserSerializer, ProgramManagerUserSerializer,
    AuditLogSerializer,
)
from users.models import (
    User, RoleConfig, UserRoleAssignment, Role,
    AdminProfile, ProfessorProfile, StudentProfile,
    HeadOfDepartmentProfile, ProgramManagerProfile,
)

import jwt
from django.conf import settings


def get_role_from_request(request):
    auth_header = request.headers.get('Authorization', '')
    if not auth_header.startswith('Bearer '):
        return None, None
    token_str = auth_header.split(' ')[1]
    try:
        payload = jwt.decode(token_str, settings.SECRET_KEY, algorithms=['HS256'])
        roles = payload.get('roles', [])
        if isinstance(roles, str):
            roles = [roles]
        user_id = payload.get('user_id')
        return roles, user_id
    except jwt.ExpiredSignatureError:
        return None, None
    except jwt.InvalidTokenError:
        return None, None


def has_role(roles, required_role):
    if not roles:
        return False
    return required_role in roles


def log_audit(request, action_type, description, related_object_type=None, related_object_id=None):
    from .models import AuditLog
    roles, user_id = get_role_from_request(request)
    user = None
    if user_id:
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            pass
    AuditLog.objects.create(
        user=user,
        action_type=action_type,
        description=description,
        related_object_type=related_object_type,
        related_object_id=related_object_id,
    )


class FiliereViewSet(viewsets.ModelViewSet):
    authentication_classes = []
    permission_classes = [AllowAny]
    queryset = Filiere.objects.all()
    serializer_class = FiliereSerializer

    def get_queryset(self):
        roles, user_id = get_role_from_request(self.request)
        if not roles or not user_id:
            return Filiere.objects.none()
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Filiere.objects.none()

        active_roles = user.get_active_roles()
        if Role.ADMIN in active_roles:
            return Filiere.objects.all()
        if Role.HEAD_OF_DEPARTMENT in active_roles:
            hod = getattr(user, 'hod_profile', None)
            if hod and hod.department_id:
                return Filiere.objects.filter(departement_id=hod.department_id)
            return Filiere.objects.none()
        if Role.PROGRAM_MANAGER in active_roles:
            pm = getattr(user, 'program_manager_profile', None)
            if pm and pm.program_id:
                return Filiere.objects.filter(id=pm.program_id)
            return Filiere.objects.none()
        if Role.PROFESSOR in active_roles:
            prof = getattr(user, 'professor_profile', None)
            if prof:
                dept_ids = prof.departements.values_list('id', flat=True)
                return Filiere.objects.filter(departement_id__in=dept_ids)
            return Filiere.objects.none()
        return Filiere.objects.none()

    def _can_write(self, user, filiere=None):
        active_roles = user.get_active_roles()
        if Role.ADMIN in active_roles:
            return True
        if Role.HEAD_OF_DEPARTMENT in active_roles:
            hod = getattr(user, 'hod_profile', None)
            if not hod or not hod.department_id:
                return False
            if filiere is not None:
                return filiere.departement_id == hod.department_id
            return True
        if Role.PROGRAM_MANAGER in active_roles:
            pm = getattr(user, 'program_manager_profile', None)
            if not pm or not pm.program_id:
                return False
            if filiere is not None:
                return filiere.id == pm.program_id
            return True
        return False

    def list(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        return super().list(request, *args, **kwargs)

    def create(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        if not self._can_write(user):
            return Response({'error': 'Acces non autorise'}, status=403)
        return super().create(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        instance = self.get_object()
        if not self._can_write(user, instance):
            return Response({'error': 'Acces non autorise a cette filiere'}, status=403)
        return super().update(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        instance = self.get_object()
        if not self._can_write(user, instance):
            return Response({'error': 'Acces non autorise a cette filiere'}, status=403)
        return super().partial_update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        instance = self.get_object()
        if not self._can_write(user, instance):
            return Response({'error': 'Acces non autorise a cette filiere'}, status=403)
        return super().destroy(request, *args, **kwargs)


class SemestreViewSet(viewsets.ModelViewSet):
    authentication_classes = []
    permission_classes = [AllowAny]
    queryset = Semestre.objects.all()
    serializer_class = SemestreSerializer

    def get_queryset(self):
        roles, user_id = get_role_from_request(self.request)
        if not roles or not user_id:
            return Semestre.objects.none()
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Semestre.objects.none()

        qs = Semestre.objects.all()

        filiere = self.request.GET.get('filiere')
        if filiere:
            qs = qs.filter(filiere_id=filiere)

        active_roles = user.get_active_roles()
        if Role.ADMIN in active_roles:
            return qs
        if Role.HEAD_OF_DEPARTMENT in active_roles:
            hod = getattr(user, 'hod_profile', None)
            if hod and hod.department_id:
                return qs.filter(filiere__departement_id=hod.department_id)
            return Semestre.objects.none()
        if Role.PROGRAM_MANAGER in active_roles:
            pm = getattr(user, 'program_manager_profile', None)
            if pm and pm.program_id:
                return qs.filter(filiere_id=pm.program_id)
            return Semestre.objects.none()
        if Role.PROFESSOR in active_roles:
            prof = getattr(user, 'professor_profile', None)
            if prof:
                dept_ids = prof.departements.values_list('id', flat=True)
                return qs.filter(filiere__departement_id__in=dept_ids)
            return Semestre.objects.none()
        return Semestre.objects.none()

    def _can_write(self, user, semestre=None):
        active_roles = user.get_active_roles()
        if Role.ADMIN in active_roles:
            return True
        if Role.HEAD_OF_DEPARTMENT in active_roles:
            hod = getattr(user, 'hod_profile', None)
            if not hod or not hod.department_id:
                return False
            if semestre is not None:
                return semestre.filiere.departement_id == hod.department_id
            return True
        if Role.PROGRAM_MANAGER in active_roles:
            pm = getattr(user, 'program_manager_profile', None)
            if not pm or not pm.program_id:
                return False
            if semestre is not None:
                return semestre.filiere_id == pm.program_id
            return True
        return False

    def create(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        if not self._can_write(user):
            return Response({'error': 'Acces non autorise'}, status=403)
        return super().create(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        instance = self.get_object()
        if not self._can_write(user, instance):
            return Response({'error': 'Acces non autorise'}, status=403)
        return super().update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        instance = self.get_object()
        if not self._can_write(user, instance):
            return Response({'error': 'Acces non autorise'}, status=403)
        return super().destroy(request, *args, **kwargs)


class ModuleViewSet(viewsets.ModelViewSet):
    authentication_classes = []
    permission_classes = [AllowAny]
    queryset = Module.objects.all()
    serializer_class = ModuleSerializer

    def get_queryset(self):
        roles, user_id = get_role_from_request(self.request)
        if not roles or not user_id:
            return Module.objects.none()
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Module.objects.none()

        qs = Module.objects.all()

        semestre = self.request.GET.get('semestre')
        if semestre:
            qs = qs.filter(semestre_id=semestre)

        active_roles = user.get_active_roles()
        if Role.ADMIN in active_roles:
            return qs
        if Role.HEAD_OF_DEPARTMENT in active_roles:
            hod = getattr(user, 'hod_profile', None)
            if hod and hod.department_id:
                return qs.filter(semestre__filiere__departement_id=hod.department_id)
            return Module.objects.none()
        if Role.PROGRAM_MANAGER in active_roles:
            pm = getattr(user, 'program_manager_profile', None)
            if pm and pm.program_id:
                return qs.filter(semestre__filiere_id=pm.program_id)
            return Module.objects.none()
        if Role.PROFESSOR in active_roles:
            prof = getattr(user, 'professor_profile', None)
            if prof:
                dept_ids = prof.departements.values_list('id', flat=True)
                return qs.filter(semestre__filiere__departement_id__in=dept_ids)
            return Module.objects.none()
        return Module.objects.none()

    def _can_write(self, user, module=None):
        active_roles = user.get_active_roles()
        if Role.ADMIN in active_roles:
            return True
        if Role.HEAD_OF_DEPARTMENT in active_roles:
            hod = getattr(user, 'hod_profile', None)
            if not hod or not hod.department_id:
                return False
            if module is not None and module.semestre:
                return module.semestre.filiere.departement_id == hod.department_id
            return True
        if Role.PROGRAM_MANAGER in active_roles:
            pm = getattr(user, 'program_manager_profile', None)
            if not pm or not pm.program_id:
                return False
            if module is not None:
                return module.semestre.filiere_id == pm.program_id
            return True
        return False

    def create(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        if not self._can_write(user):
            return Response({'error': 'Acces non autorise'}, status=403)
        return super().create(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        instance = self.get_object()
        if not self._can_write(user, instance):
            return Response({'error': 'Acces non autorise'}, status=403)
        return super().update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        instance = self.get_object()
        if not self._can_write(user, instance):
            return Response({'error': 'Acces non autorise'}, status=403)
        return super().destroy(request, *args, **kwargs)


class DepartementViewSet(viewsets.ModelViewSet):
    authentication_classes = []
    permission_classes = [AllowAny]
    queryset = Departement.objects.all()
    serializer_class = DepartementSerializer

    def get_queryset(self):
        roles, user_id = get_role_from_request(self.request)
        if not roles or not user_id:
            return Departement.objects.none()
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Departement.objects.none()

        active_roles = user.get_active_roles()
        if Role.ADMIN in active_roles:
            return Departement.objects.all()
        if Role.HEAD_OF_DEPARTMENT in active_roles:
            hod = getattr(user, 'hod_profile', None)
            if hod and hod.department_id:
                return Departement.objects.filter(id=hod.department_id)
            return Departement.objects.none()
        return Departement.objects.none()

    def _can_write(self, user, dept=None):
        active_roles = user.get_active_roles()
        if Role.ADMIN in active_roles:
            return True
        if Role.HEAD_OF_DEPARTMENT in active_roles:
            hod = getattr(user, 'hod_profile', None)
            if not hod or not hod.department_id:
                return False
            if dept is not None:
                return dept.id == hod.department_id
            return True
        return False

    def list(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        return super().list(request, *args, **kwargs)

    def create(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        if not self._can_write(user):
            return Response({'error': 'Acces non autorise'}, status=403)
        return super().create(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        instance = self.get_object()
        if not self._can_write(user, instance):
            return Response({'error': 'Acces non autorise a ce departement'}, status=403)
        return super().update(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        instance = self.get_object()
        if not self._can_write(user, instance):
            return Response({'error': 'Acces non autorise a ce departement'}, status=403)
        return super().partial_update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        instance = self.get_object()
        if not self._can_write(user, instance):
            return Response({'error': 'Acces non autorise a ce departement'}, status=403)
        return super().destroy(request, *args, **kwargs)


class SalleViewSet(viewsets.ModelViewSet):
    authentication_classes = []
    permission_classes = [AllowAny]
    queryset = Salle.objects.all()
    serializer_class = SalleSerializer


class AdminListViewSet(viewsets.ModelViewSet):
    authentication_classes = []
    permission_classes = [AllowAny]

    def get_queryset(self):
        return User.objects.filter(
            role_assignments__role__codename='admin',
            role_assignments__is_active=True
        )

    def get_serializer_class(self):
        return ProfessorUserSerializer

    def list(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not has_role(roles, 'admin'):
            return Response([], status=200)
        return super().list(request, *args, **kwargs)

    def create(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not has_role(roles, 'admin'):
            return Response({'error': 'Reserve aux administrateurs'}, status=403)
        return self._create_user(request, Role.ADMIN)

    def update(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not has_role(roles, 'admin'):
            return Response({'error': 'Reserve aux administrateurs'}, status=403)
        return self._update_user(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not has_role(roles, 'admin'):
            return Response({'error': 'Reserve aux administrateurs'}, status=403)
        return self._update_user(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not has_role(roles, 'admin'):
            return Response({'error': 'Reserve aux administrateurs'}, status=403)
        return super().destroy(request, *args, **kwargs)

    def _create_user(self, request, role):
        email = request.data.get('email')
        password = request.data.get('password')
        nom = request.data.get('last_name', '')
        prenom = request.data.get('first_name', '')

        if not email or not password:
            return Response({'error': 'email et password obligatoires'}, status=400)

        if User.objects.filter(email=email).exists():
            return Response({'error': 'Email déjà utilisé'}, status=400)

        user = User.objects.create_user(
            email=email, password=password,
            first_name=prenom, last_name=nom,
        )

        from users.views import create_role_assignments, create_profile_for_role
        all_roles = [role]
        extras = [{'role': role, 'data': {}}]

        if request.data.get('also_professor'):
            all_roles.append(Role.PROFESSOR)
            extras.append({'role': Role.PROFESSOR, 'data': {
                'specialization': request.data.get('specialization', ''),
                'department_ids': request.data.get('department_ids', []),
            }})

        if request.data.get('also_head_of_department'):
            all_roles.append(Role.HEAD_OF_DEPARTMENT)
            extras.append({'role': Role.HEAD_OF_DEPARTMENT, 'data': {
                'department_id': request.data.get('department_id'),
            }})

        if request.data.get('also_program_manager'):
            all_roles.append(Role.PROGRAM_MANAGER)
            extras.append({'role': Role.PROGRAM_MANAGER, 'data': {
                'program_id': request.data.get('program_id'),
            }})

        create_role_assignments(user, all_roles)
        for extra in extras:
            create_profile_for_role(user, extra['role'], extra['data'])

        serializer = self.get_serializer(user)
        return Response(serializer.data, status=201)

    def _update_user(self, request, *args, **kwargs):
        kwargs.pop('partial', False)
        instance = self.get_object()
        email = request.data.get('email', instance.email)
        nom = request.data.get('last_name', instance.last_name)
        prenom = request.data.get('first_name', instance.first_name)
        password = request.data.get('password')

        instance.email = email
        instance.last_name = nom
        instance.first_name = prenom
        if password:
            instance.set_password(password)
        instance.save()

        serializer = self.get_serializer(instance)
        return Response(serializer.data)


class StudentListViewSet(viewsets.ModelViewSet):
    authentication_classes = []
    permission_classes = [AllowAny]

    def get_queryset(self):
        qs = User.objects.filter(
            role_assignments__role__codename='student',
            role_assignments__is_active=True
        )
        roles, user_id = get_role_from_request(self.request)
        if not roles or not user_id:
            return qs.none()
        if has_role(roles, 'admin'):
            return qs
        if has_role(roles, 'head_of_department'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                return qs.none()
            hod = getattr(user, 'hod_profile', None)
            if hod and hod.department_id:
                return qs.filter(
                    student_profile__program__departement_id=hod.department_id
                )
            return qs.none()
        if has_role(roles, 'program_manager'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                return qs.none()
            pm = getattr(user, 'program_manager_profile', None)
            if pm and pm.program_id:
                return qs.filter(student_profile__program_id=pm.program_id)
            return qs.none()
        if has_role(roles, 'professor'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                return qs.none()
            prof = getattr(user, 'professor_profile', None)
            if prof:
                dept_ids = prof.departements.values_list('id', flat=True)
                return qs.filter(student_profile__program__departement_id__in=dept_ids)
            return qs.none()
        return qs.none()

    def get_serializer_class(self):
        return StudentUserSerializer

    def list(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not roles:
            return Response([], status=200)
        if has_role(roles, 'head_of_department') or has_role(roles, 'admin') or has_role(roles, 'program_manager') or has_role(roles, 'professor'):
            return super().list(request, *args, **kwargs)
        return Response([], status=200)

    def _get_user_and_roles(self, request):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return None, None
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return None, None
        return user, roles

    def _can_access_student(self, user, roles, student_user=None):
        if has_role(roles, 'admin') or Role.ADMIN in user.get_active_roles():
            return True
        if has_role(roles, 'head_of_department'):
            hod = getattr(user, 'hod_profile', None)
            if not hod or not hod.department_id:
                return False
            if student_user:
                sp = getattr(student_user, 'student_profile', None)
                if sp and sp.program:
                    return sp.program.departement_id == hod.department_id
                return False
            return True
        if has_role(roles, 'program_manager'):
            pm = getattr(user, 'program_manager_profile', None)
            if not pm or not pm.program_id:
                return False
            if student_user:
                sp = getattr(student_user, 'student_profile', None)
                if sp:
                    return sp.program_id == pm.program_id
                return False
            return True
        return False

    def create(self, request, *args, **kwargs):
        user, roles = self._get_user_and_roles(request)
        filiere_id = request.data.get('program_id') or request.data.get('filiere_id')
        debug_info = {
            'user_found': user is not None,
            'token_roles': roles,
            'db_roles': user.get_active_roles() if user else [],
            'user_email': getattr(user, 'email', None) if user else None,
            'user_id': user.id if user else None,
            'filiere_id': filiere_id,
            'can_access': self._can_access_student(user, roles) if user else False,
            'has_role_admin_in_token': has_role(roles, 'admin') if roles else False,
            'has_role_admin_in_db': Role.ADMIN in (user.get_active_roles() if user else []),
        }
        if not user or not self._can_access_student(user, roles):
            logger.error(f"403 STUDENT DEBUG: {debug_info}")
            return Response({'debug': debug_info}, status=403)
        if has_role(roles, 'program_manager'):
            pm = getattr(user, 'program_manager_profile', None)
            if not pm or not pm.program_id:
                return Response({'error': 'Acces non autorise'}, status=403)
            if not filiere_id or int(filiere_id) != pm.program_id:
                return Response({'error': 'Filiere hors de votre ressort'}, status=403)
        return self._create_user(request, Role.STUDENT)

    def update(self, request, *args, **kwargs):
        user, roles = self._get_user_and_roles(request)
        if not user:
            return Response({'error': 'Authentification requise'}, status=401)
        instance = self.get_object()
        if not self._can_access_student(user, roles, instance):
            return Response({'error': 'Acces non autorise a cet etudiant'}, status=403)
        return self._update_user(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        user, roles = self._get_user_and_roles(request)
        if not user:
            return Response({'error': 'Authentification requise'}, status=401)
        instance = self.get_object()
        if not self._can_access_student(user, roles, instance):
            return Response({'error': 'Acces non autorise a cet etudiant'}, status=403)
        return self._update_user(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        user, roles = self._get_user_and_roles(request)
        if not user:
            return Response({'error': 'Authentification requise'}, status=401)
        instance = self.get_object()
        if not self._can_access_student(user, roles, instance):
            return Response({'error': 'Acces non autorise a cet etudiant'}, status=403)
        info = f"{instance.first_name} {instance.last_name} ({instance.email})"
        result = super().destroy(request, *args, **kwargs)
        log_audit(request, 'DELETE_STUDENT', f"Suppression étudiant: {info}", 'User', instance.id)
        return result

    def _create_user(self, request, role):
        email = request.data.get('email')
        password = request.data.get('password')
        nom = request.data.get('last_name', '')
        prenom = request.data.get('first_name', '')
        cie = request.data.get('CIE', request.data.get('code_massar', ''))
        program_id = request.data.get('program_id') or request.data.get('filiere_id')
        module_ids = request.data.get('module_ids', [])

        if not email or not password:
            return Response({'error': 'email et password obligatoires'}, status=400)

        if User.objects.filter(email=email).exists():
            return Response({'error': 'Email déjà utilisé'}, status=400)

        user = User.objects.create_user(
            email=email, password=password,
            first_name=prenom, last_name=nom,
        )

        from users.views import create_role_assignments, create_profile_for_role
        create_role_assignments(user, [role])
        create_profile_for_role(user, role, {
            'CIE': cie,
            'program_id': program_id,
        })

        if module_ids:
            profile = getattr(user, 'student_profile', None)
            if profile:
                profile.modules.set(module_ids)

        serializer = self.get_serializer(user)
        log_audit(request, 'CREATE_STUDENT', f"Ajout étudiant: {prenom} {nom} ({email})", 'User', user.id)
        return Response(serializer.data, status=201)

    def _update_user(self, request, *args, **kwargs):
        kwargs.pop('partial', False)
        instance = self.get_object()
        email = request.data.get('email', instance.email)
        nom = request.data.get('last_name', instance.last_name)
        prenom = request.data.get('first_name', instance.first_name)
        password = request.data.get('password')

        old_email = instance.email
        instance.email = email
        instance.last_name = nom
        instance.first_name = prenom
        if password:
            instance.set_password(password)
        instance.save()

        program_id = request.data.get('program_id') or request.data.get('filiere_id')
        if program_id is not None:
            profile = getattr(instance, 'student_profile', None)
            if profile:
                profile.program_id = program_id
                profile.save()

        module_ids = request.data.get('module_ids')
        if module_ids is not None:
            profile = getattr(instance, 'student_profile', None)
            if profile:
                profile.modules.set(module_ids)

        serializer = self.get_serializer(instance)
        log_audit(request, 'UPDATE_STUDENT', f"Modification étudiant: {prenom} {nom} ({old_email})", 'User', instance.id)
        return Response(serializer.data)


class ProfessorListViewSet(viewsets.ModelViewSet):
    authentication_classes = []
    permission_classes = [AllowAny]

    def get_queryset(self):
        qs = User.objects.filter(
            role_assignments__role__codename='professor',
            role_assignments__is_active=True
        )
        roles, user_id = get_role_from_request(self.request)
        if not roles or not user_id:
            return qs.none()
        if has_role(roles, 'admin'):
            return qs
        if has_role(roles, 'head_of_department'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                return qs.none()
            hod = getattr(user, 'hod_profile', None)
            if hod and hod.department_id:
                return qs.filter(
                    professor_profile__departements=hod.department_id
                )
            return qs.none()
        if has_role(roles, 'program_manager'):
            return qs
        return qs.none()

    def get_serializer_class(self):
        return ProfessorUserSerializer

    def _get_user_and_dept(self, request):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return None, None
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return None, None
        hod_dept_id = None
        if has_role(roles, 'head_of_department'):
            hod = getattr(user, 'hod_profile', None)
            if hod:
                hod_dept_id = hod.department_id
        return user, hod_dept_id

    def _can_manage_professor(self, user, hod_dept_id, professor_user=None):
        if Role.ADMIN in user.get_active_roles():
            return True
        if hod_dept_id:
            if professor_user:
                prof_profile = getattr(professor_user, 'professor_profile', None)
                if not prof_profile:
                    return False
                return prof_profile.departements.filter(id=hod_dept_id).exists()
            return True
        return False

    def list(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not roles:
            return Response([], status=200)
        if has_role(roles, 'head_of_department') or has_role(roles, 'admin') or has_role(roles, 'program_manager'):
            return super().list(request, *args, **kwargs)
        return Response([], status=200)

    def create(self, request, *args, **kwargs):
        user, hod_dept_id = self._get_user_and_dept(request)
        if not user or not self._can_manage_professor(user, hod_dept_id):
            return Response({'error': 'Acces non autorise'}, status=403)
        return self._create_user(request, Role.PROFESSOR, hod_dept_id)

    def update(self, request, *args, **kwargs):
        user, hod_dept_id = self._get_user_and_dept(request)
        if not user:
            return Response({'error': 'Authentification requise'}, status=401)
        instance = self.get_object()
        if not self._can_manage_professor(user, hod_dept_id, instance):
            return Response({'error': 'Acces non autorise a cet enseignant'}, status=403)
        return self._update_user(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        user, hod_dept_id = self._get_user_and_dept(request)
        if not user:
            return Response({'error': 'Authentification requise'}, status=401)
        instance = self.get_object()
        if not self._can_manage_professor(user, hod_dept_id, instance):
            return Response({'error': 'Acces non autorise a cet enseignant'}, status=403)
        return self._update_user(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        user, hod_dept_id = self._get_user_and_dept(request)
        if not user:
            return Response({'error': 'Authentification requise'}, status=401)
        instance = self.get_object()
        if not self._can_manage_professor(user, hod_dept_id, instance):
            return Response({'error': 'Acces non autorise a cet enseignant'}, status=403)
        info = f"{instance.first_name} {instance.last_name} ({instance.email})"
        result = super().destroy(request, *args, **kwargs)
        log_audit(request, 'DELETE_TEACHER', f"Suppression enseignant: {info}", 'User', instance.id)
        return result

    def _create_user(self, request, role, hod_dept_id=None):
        email = request.data.get('email')
        password = request.data.get('password')
        nom = request.data.get('last_name', '')
        prenom = request.data.get('first_name', '')

        if not email or not password:
            return Response({'error': 'email et password obligatoires'}, status=400)

        if User.objects.filter(email=email).exists():
            return Response({'error': 'Email déjà utilisé'}, status=400)

        user = User.objects.create_user(
            email=email, password=password,
            first_name=prenom, last_name=nom,
        )

        from users.views import create_role_assignments, create_profile_for_role
        extra_data = {
            'specialization': request.data.get('specialization', ''),
        }
        department_ids = request.data.get('department_ids')
        if department_ids:
            extra_data['department_ids'] = department_ids
        elif hod_dept_id:
            extra_data['department_ids'] = [hod_dept_id]

        create_role_assignments(user, [role])
        create_profile_for_role(user, role, extra_data)

        serializer = self.get_serializer(user)
        log_audit(request, 'CREATE_TEACHER', f"Ajout enseignant: {prenom} {nom} ({email})", 'User', user.id)
        return Response(serializer.data, status=201)

    def _update_user(self, request, *args, **kwargs):
        kwargs.pop('partial', False)
        instance = self.get_object()
        email = request.data.get('email', instance.email)
        nom = request.data.get('last_name', instance.last_name)
        prenom = request.data.get('first_name', instance.first_name)
        password = request.data.get('password')

        old_email = instance.email
        instance.email = email
        instance.last_name = nom
        instance.first_name = prenom
        if password:
            instance.set_password(password)
        instance.save()

        department_ids = request.data.get('department_ids')
        if department_ids is not None:
            profile = getattr(instance, 'professor_profile', None)
            if profile:
                profile.departements.set(department_ids)

        serializer = self.get_serializer(instance)
        log_audit(request, 'UPDATE_TEACHER', f"Modification enseignant: {prenom} {nom} ({old_email})", 'User', instance.id)
        return Response(serializer.data)


class HeadOfDepartmentListViewSet(viewsets.ModelViewSet):
    authentication_classes = []
    permission_classes = [AllowAny]

    def get_queryset(self):
        return User.objects.filter(
            role_assignments__role__codename='head_of_department',
            role_assignments__is_active=True
        )

    def get_serializer_class(self):
        return HeadOfDepartmentUserSerializer

    def list(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not has_role(roles, 'admin'):
            return Response([], status=200)
        return super().list(request, *args, **kwargs)

    def create(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not has_role(roles, 'admin'):
            return Response({'error': 'Reserve aux administrateurs'}, status=403)
        return self._create_user(request, Role.HEAD_OF_DEPARTMENT)

    def update(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not has_role(roles, 'admin'):
            return Response({'error': 'Reserve aux administrateurs'}, status=403)
        return self._update_user(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not has_role(roles, 'admin'):
            return Response({'error': 'Reserve aux administrateurs'}, status=403)
        return self._update_user(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not has_role(roles, 'admin'):
            return Response({'error': 'Reserve aux administrateurs'}, status=403)
        return super().destroy(request, *args, **kwargs)

    def _create_user(self, request, role):
        email = request.data.get('email')
        password = request.data.get('password')
        nom = request.data.get('last_name', '')
        prenom = request.data.get('first_name', '')

        if not email or not password:
            return Response({'error': 'email et password obligatoires'}, status=400)

        if User.objects.filter(email=email).exists():
            return Response({'error': 'Email déjà utilisé'}, status=400)

        user = User.objects.create_user(
            email=email, password=password,
            first_name=prenom, last_name=nom,
        )

        from users.views import create_role_assignments, create_profile_for_role
        all_roles = [role]
        extras = [{'role': role, 'data': {
            'department_id': request.data.get('department_id'),
        }}]

        if request.data.get('also_professor'):
            all_roles.append(Role.PROFESSOR)
            extras.append({'role': Role.PROFESSOR, 'data': {
                'specialization': request.data.get('specialization', ''),
            }})

        if request.data.get('also_program_manager'):
            all_roles.append(Role.PROGRAM_MANAGER)
            extras.append({'role': Role.PROGRAM_MANAGER, 'data': {
                'program_id': request.data.get('program_id'),
            }})

        create_role_assignments(user, all_roles)
        for extra in extras:
            create_profile_for_role(user, extra['role'], extra['data'])

        serializer = self.get_serializer(user)
        return Response(serializer.data, status=201)

    def _update_user(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        email = request.data.get('email', instance.email)
        nom = request.data.get('last_name', instance.last_name)
        prenom = request.data.get('first_name', instance.first_name)
        password = request.data.get('password')

        instance.email = email
        instance.last_name = nom
        instance.first_name = prenom
        if password:
            instance.set_password(password)
        instance.save()

        department_id = request.data.get('department_id')
        if department_id is not None:
            profile = getattr(instance, 'hod_profile', None)
            if profile:
                profile.department_id = department_id
                profile.save()

        serializer = self.get_serializer(instance)
        return Response(serializer.data)


class ProgramManagerListViewSet(viewsets.ModelViewSet):
    authentication_classes = []
    permission_classes = [AllowAny]

    def get_queryset(self):
        return User.objects.filter(
            role_assignments__role__codename='program_manager',
            role_assignments__is_active=True
        )

    def get_serializer_class(self):
        return ProgramManagerUserSerializer

    def list(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not has_role(roles, 'admin'):
            return Response([], status=200)
        return super().list(request, *args, **kwargs)

    def create(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not has_role(roles, 'admin'):
            return Response({'error': 'Reserve aux administrateurs'}, status=403)
        return self._create_user(request, Role.PROGRAM_MANAGER)

    def update(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not has_role(roles, 'admin'):
            return Response({'error': 'Reserve aux administrateurs'}, status=403)
        return self._update_user(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not has_role(roles, 'admin'):
            return Response({'error': 'Reserve aux administrateurs'}, status=403)
        return self._update_user(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not has_role(roles, 'admin'):
            return Response({'error': 'Reserve aux administrateurs'}, status=403)
        return super().destroy(request, *args, **kwargs)

    def _create_user(self, request, role):
        email = request.data.get('email')
        password = request.data.get('password')
        nom = request.data.get('last_name', '')
        prenom = request.data.get('first_name', '')

        if not email or not password:
            return Response({'error': 'email et password obligatoires'}, status=400)

        if User.objects.filter(email=email).exists():
            return Response({'error': 'Email déjà utilisé'}, status=400)

        user = User.objects.create_user(
            email=email, password=password,
            first_name=prenom, last_name=nom,
        )

        from users.views import create_role_assignments, create_profile_for_role
        all_roles = [role]
        extras = [{'role': role, 'data': {
            'program_name': request.data.get('program_name', ''),
            'program_id': request.data.get('program_id'),
        }}]

        if request.data.get('also_professor'):
            all_roles.append(Role.PROFESSOR)
            extras.append({'role': Role.PROFESSOR, 'data': {
                'specialization': request.data.get('specialization', ''),
            }})

        create_role_assignments(user, all_roles)
        for extra in extras:
            create_profile_for_role(user, extra['role'], extra['data'])

        serializer = self.get_serializer(user)
        return Response(serializer.data, status=201)

    def _update_user(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        email = request.data.get('email', instance.email)
        nom = request.data.get('last_name', instance.last_name)
        prenom = request.data.get('first_name', instance.first_name)
        password = request.data.get('password')

        instance.email = email
        instance.last_name = nom
        instance.first_name = prenom
        if password:
            instance.set_password(password)
        instance.save()

        program_id = request.data.get('program_id')
        if program_id is not None:
            profile = getattr(instance, 'program_manager_profile', None)
            if profile:
                profile.program_id = program_id
                profile.save()

        serializer = self.get_serializer(instance)
        return Response(serializer.data)


class PlanificationViewSet(viewsets.ModelViewSet):
    authentication_classes = []
    permission_classes = [AllowAny]
    queryset = Planification.objects.all()
    serializer_class = PlanificationSerializer

    def get_queryset(self):
        roles, user_id = get_role_from_request(self.request)
        if not roles or not user_id:
            return Planification.objects.none()
        if has_role(roles, 'admin'):
            return Planification.objects.all()
        if has_role(roles, 'head_of_department'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                return Planification.objects.none()
            hod = getattr(user, 'hod_profile', None)
            if hod and hod.department_id:
                return Planification.objects.filter(
                    filiere__departement_id=hod.department_id
                )
            return Planification.objects.none()
        if has_role(roles, 'program_manager'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                return Planification.objects.none()
            pm = getattr(user, 'program_manager_profile', None)
            if pm and pm.program_id:
                return Planification.objects.filter(
                    filiere_id=pm.program_id
                )
            return Planification.objects.none()
        if has_role(roles, 'professor'):
            return Planification.objects.filter(professor_id=user_id)
        return Planification.objects.none()

    def _can_manage(self, roles, user_id, planification=None):
        if has_role(roles, 'admin'):
            return True
        if has_role(roles, 'head_of_department'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                return False
            hod = getattr(user, 'hod_profile', None)
            if not hod or not hod.department_id:
                return False
            if planification is not None:
                return planification.filiere.departement_id == hod.department_id
            return True
        if has_role(roles, 'program_manager'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                return False
            pm = getattr(user, 'program_manager_profile', None)
            if not pm or not pm.program_id:
                return False
            if planification is not None:
                return planification.filiere_id == pm.program_id
            return True
        if has_role(roles, 'professor'):
            if planification is not None:
                return planification.professor.id == user_id
            return True
        return False

    def list(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not roles:
            return Response([], status=200)
        return super().list(request, *args, **kwargs)

    def create(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        if not self._can_manage(roles, user_id):
            return Response({'error': 'Acces non autorise'}, status=403)
        filiere_id = request.data.get('filiere')
        if not filiere_id:
            return Response({'error': 'Filiere requise'}, status=400)
        if has_role(roles, 'head_of_department') and not has_role(roles, 'admin'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                return Response({'error': 'Utilisateur introuvable'}, status=404)
            hod = getattr(user, 'hod_profile', None)
            if not hod or not hod.department_id:
                return Response({'error': 'Acces non autorise'}, status=403)
            if not Filiere.objects.filter(id=filiere_id, departement_id=hod.department_id).exists():
                return Response({'error': 'Filiere hors de votre departement'}, status=403)
        if has_role(roles, 'program_manager') and not has_role(roles, 'admin'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                return Response({'error': 'Utilisateur introuvable'}, status=404)
            pm = getattr(user, 'program_manager_profile', None)
            if not pm or not pm.program_id:
                return Response({'error': 'Acces non autorise'}, status=403)
            if int(filiere_id) != pm.program_id:
                return Response({'error': 'Filiere hors de votre ressort'}, status=403)
        semestre_id = request.data.get('semestre')
        module_ref_id = request.data.get('module_ref')
        if semestre_id and filiere_id:
            if not Semestre.objects.filter(id=semestre_id, filiere_id=filiere_id).exists():
                return Response({'error': 'Semestre invalide pour cette filiere'}, status=400)
        if module_ref_id and semestre_id:
            if not Module.objects.filter(id=module_ref_id, semestre_id=semestre_id).exists():
                return Response({'error': 'Module invalide pour ce semestre'}, status=400)
        return super().create(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        instance = self.get_object()
        if not self._can_manage(roles, user_id, instance):
            return Response({'error': 'Acces non autorise'}, status=403)
        return super().update(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        instance = self.get_object()
        if not self._can_manage(roles, user_id, instance):
            return Response({'error': 'Acces non autorise'}, status=403)
        return super().partial_update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not roles or not user_id:
            return Response({'error': 'Authentification requise'}, status=401)
        instance = self.get_object()
        if not self._can_manage(roles, user_id, instance):
            return Response({'error': 'Acces non autorise'}, status=403)
        return super().destroy(request, *args, **kwargs)


class SeanceQRViewSet(viewsets.ModelViewSet):
    authentication_classes = []
    permission_classes = [AllowAny]
    queryset = SeanceQR.objects.all()
    serializer_class = SeanceQRSerializer

    def get_queryset(self):
        roles, user_id = get_role_from_request(self.request)
        if not roles or not user_id:
            return SeanceQR.objects.none()
        if has_role(roles, 'admin'):
            return SeanceQR.objects.all()
        if has_role(roles, 'head_of_department'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                return SeanceQR.objects.none()
            hod = getattr(user, 'hod_profile', None)
            if hod and hod.department_id:
                return SeanceQR.objects.filter(
                    planification__filiere__departement_id=hod.department_id
                )
            return SeanceQR.objects.none()
        if has_role(roles, 'program_manager'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                return SeanceQR.objects.none()
            pm = getattr(user, 'program_manager_profile', None)
            if pm and pm.program_id:
                return SeanceQR.objects.filter(
                    planification__filiere_id=pm.program_id
                )
            return SeanceQR.objects.none()
        if has_role(roles, 'professor'):
            return SeanceQR.objects.filter(planification__professor_id=user_id)
        return SeanceQR.objects.none()

    def list(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not roles:
            return Response([], status=200)
        return super().list(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not has_role(roles, 'professor'):
            return Response({'error': 'Reserve aux enseignants'}, status=403)
        seance = self.get_object()
        if seance.planification.professor.id != user_id:
            return Response({'error': 'Acces non autorise a cette seance'}, status=403)
        return super().destroy(request, *args, **kwargs)


class QRCodeViewSet(viewsets.ModelViewSet):
    authentication_classes = []
    permission_classes = [AllowAny]
    queryset = QRCode.objects.all()
    serializer_class = QRCodeSerializer


class PresenceViewSet(viewsets.ModelViewSet):
    authentication_classes = []
    permission_classes = [AllowAny]
    queryset = Presence.objects.all()
    serializer_class = PresenceSerializer

    def get_queryset(self):
        roles, user_id = get_role_from_request(self.request)
        if not roles or not user_id:
            return Presence.objects.none()
        if has_role(roles, 'admin'):
            return Presence.objects.all()
        if has_role(roles, 'head_of_department'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                return Presence.objects.none()
            hod = getattr(user, 'hod_profile', None)
            if hod and hod.department_id:
                return Presence.objects.filter(
                    seance__planification__filiere__departement_id=hod.department_id
                )
            return Presence.objects.none()
        if has_role(roles, 'program_manager'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                return Presence.objects.none()
            pm = getattr(user, 'program_manager_profile', None)
            if pm and pm.program_id:
                return Presence.objects.filter(
                    seance__planification__filiere_id=pm.program_id
                )
            return Presence.objects.none()
        if has_role(roles, 'professor'):
            return Presence.objects.filter(
                seance__planification__professor_id=user_id
            )
        return Presence.objects.none()

    def list(self, request, *args, **kwargs):
        roles, _ = get_role_from_request(request)
        if not roles:
            return Response([], status=200)
        return super().list(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        roles, user_id = get_role_from_request(request)
        if not has_role(roles, 'professor'):
            return Response({'error': 'Reserve aux enseignants'}, status=403)
        presence = self.get_object()
        if presence.seance.planification.professor.id != user_id:
            return Response({'error': 'Acces non autorise a cette presence'}, status=403)
        return super().destroy(request, *args, **kwargs)


@api_view(['POST'])
@authentication_classes([])
@permission_classes([AllowAny])
def generer_seance_qr(request):
    planification_id = request.data.get('planification_id')
    latitude_prof    = request.data.get('latitude_prof')
    longitude_prof   = request.data.get('longitude_prof')

    if not latitude_prof or not longitude_prof:
        return Response({'error': 'latitude_prof et longitude_prof obligatoires'}, status=400)

    if not planification_id:
        return Response({'error': 'planification_id est requis'}, status=status.HTTP_400_BAD_REQUEST)

    try:
        planification = Planification.objects.get(id=planification_id)
    except Planification.DoesNotExist:
        return Response({'error': 'Planification introuvable'}, status=status.HTTP_404_NOT_FOUND)

    auth_header = request.headers.get('Authorization', '')
    try:
        token = auth_header.split(' ')[1]
        decoded = AT(token)
        roles = decoded.get('roles', [])
        if isinstance(roles, str):
            roles = [roles]
        user_id = decoded.get('user_id')
    except Exception:
        return Response({'error': 'Token invalide'}, status=status.HTTP_401_UNAUTHORIZED)

    if not has_role(roles, 'professor'):
        return Response({'error': 'Acces reserve aux enseignants'}, status=status.HTTP_403_FORBIDDEN)

    try:
        professor = User.objects.get(id=user_id)
    except User.DoesNotExist:
        return Response({'error': 'Enseignant introuvable'}, status=status.HTTP_404_NOT_FOUND)

    if planification.professor != professor:
        return Response({'error': 'Cette planification ne vous appartient pas'}, status=status.HTTP_403_FORBIDDEN)

    seance = SeanceQR.objects.create(
        planification=planification,
        latitude_prof=latitude_prof,
        longitude_prof=longitude_prof,
    )
    qrcode = QRCode.objects.create(seance=seance)

    expire_dans = int((qrcode.cree_a + timedelta(minutes=qrcode.expiration) - timezone.now()).total_seconds())

    log_audit(request, 'CREATE_SEANCE', f"Création séance #{seance.id} - {planification.module} ({planification.jour})", 'SeanceQR', seance.id)

    return Response({
        'seance_id': seance.id,
        'token': str(qrcode.token),
        'expire_dans': expire_dans
    }, status=status.HTTP_201_CREATED)


@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def get_qrcode_seance(request, seance_id):
    auth_header = request.headers.get('Authorization', '')
    try:
        token = auth_header.split(' ')[1]
        decoded = AT(token)
        roles = decoded.get('roles', [])
        if isinstance(roles, str):
            roles = [roles]
        user_id = decoded.get('user_id')
    except Exception:
        return Response({'error': 'Token invalide'}, status=status.HTTP_401_UNAUTHORIZED)

    if not has_role(roles, 'professor'):
        return Response({'error': 'Acces reserve aux enseignants'}, status=status.HTTP_403_FORBIDDEN)

    try:
        professor = User.objects.get(id=user_id)
    except User.DoesNotExist:
        return Response({'error': 'Enseignant introuvable'}, status=status.HTTP_404_NOT_FOUND)

    try:
        seance = SeanceQR.objects.get(id=seance_id)
    except SeanceQR.DoesNotExist:
        return Response({'error': 'Seance introuvable'}, status=status.HTTP_404_NOT_FOUND)

    if seance.planification.professor != professor:
        return Response({'error': 'Cette seance ne vous appartient pas'}, status=status.HTTP_403_FORBIDDEN)

    qrcode = QRCode.objects.filter(seance=seance).order_by('-cree_a').first()

    if not qrcode or not qrcode.is_valid():
        qrcode = QRCode.objects.create(seance=seance)

    expire_dans = int((qrcode.cree_a + timedelta(minutes=qrcode.expiration) - timezone.now()).total_seconds())

    return Response({
        'token': str(qrcode.token),
        'seance_id': seance.id,
        'expire_dans': max(expire_dans, 0)
    })


JOURS_MAP = {
    0: 'lundi',
    1: 'mardi',
    2: 'mercredi',
    3: 'jeudi',
    4: 'vendredi',
    5: 'samedi',
    6: 'dimanche',
}


@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def planifications_du_jour(request):
    auth_header = request.headers.get('Authorization', '')
    if not auth_header.startswith('Bearer '):
        return Response({'error': 'Token manquant'}, status=status.HTTP_401_UNAUTHORIZED)

    token_str = auth_header.split(' ')[1]
    try:
        token = AT(token_str)
        roles = token.get('roles', [])
        if isinstance(roles, str):
            roles = [roles]
        user_id = token.get('user_id')
    except Exception:
        return Response({'error': 'Token invalide'}, status=status.HTTP_401_UNAUTHORIZED)

    if not has_role(roles, 'professor'):
        return Response({'error': 'Acces reserve aux enseignants'}, status=status.HTTP_403_FORBIDDEN)

    try:
        professor = User.objects.get(id=user_id)
    except User.DoesNotExist:
        return Response({'error': 'Enseignant introuvable'}, status=status.HTTP_404_NOT_FOUND)

    jour_actuel = JOURS_MAP[timezone.now().weekday()]

    planifications = Planification.objects.filter(
        professor=professor,
        jour=jour_actuel
    ).select_related('salle', 'filiere')

    data = []
    for p in planifications:
        data.append({
            'id':          p.id,
            'salle':       str(p.salle),
            'filiere':     str(p.filiere),
            'module':      p.module,
            'heure_debut': str(p.heure_debut),
            'heure_fin':   str(p.heure_fin),
            'jour':        p.jour,
        })

    return Response({
        'jour':           jour_actuel,
        'date':           str(timezone.localtime().date()),
        'planifications': data,
        'total':          len(data),
    })


@api_view(['POST'])
@authentication_classes([])
@permission_classes([AllowAny])
def verifier_etudiant(request):
    data = request.data
    cie    = data.get('CIE') or data.get('massar') or data.get('code_massar')
    email  = data.get('email') or data.get('email_etudiant')

    if not cie and not email:
        return Response(
            {'error': 'CIE (Massar) ou email obligatoire pour la verification'},
            status=400
        )

    profile = None
    if cie:
        try:
            profile = StudentProfile.objects.get(CIE=cie)
        except StudentProfile.DoesNotExist:
            pass

    if not profile and email:
        try:
            user = User.objects.get(email=email)
            profile = getattr(user, 'student_profile', None)
        except User.DoesNotExist:
            pass

    if not profile:
        return Response({
            'found': False,
            'message': 'Etudiant introuvable dans le systeme',
            'recherche': {
                'CIE': cie,
                'email': email,
            }
        }, status=404)

    user = profile.user
    return Response({
        'found': True,
        'message': f'Etudiant trouve : {user.first_name} {user.last_name}',
        'etudiant': {
            'id':       user.id,
            'nom':      user.last_name,
            'prenom':   user.first_name,
            'email':    user.email,
            'CIE':      profile.CIE,
            'filiere':  profile.program.nom_filiere if profile.program else None,
        }
    }, status=200)


@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def liste_presences_seance(request, seance_id):
    auth_header = request.headers.get('Authorization', '')
    if not auth_header.startswith('Bearer '):
        return Response({'error': 'Token manquant'}, status=401)

    token_str = auth_header.split(' ')[1]
    try:
        payload = jwt.decode(token_str, settings.SECRET_KEY, algorithms=['HS256'])
    except jwt.ExpiredSignatureError:
        return Response({'error': 'Token expire'}, status=401)
    except jwt.InvalidTokenError:
        return Response({'error': 'Token invalide'}, status=401)

    roles = payload.get('roles', [])
    if isinstance(roles, str):
        roles = [roles]
    if not has_role(roles, 'professor'):
        return Response({'error': 'Acces reserve aux enseignants'}, status=403)

    user_id = payload.get('user_id')

    try:
        seance = SeanceQR.objects.get(id=seance_id)
    except SeanceQR.DoesNotExist:
        return Response({'error': 'Seance introuvable'}, status=404)

    if seance.planification.professor.id != user_id:
        return Response({'error': 'Acces non autorise a cette seance'}, status=403)

    presences = Presence.objects.filter(seance=seance).select_related('student')

    data = {
        "seance_id": seance.id,
        "date": str(seance.date),
        "planification": {
            "filiere": seance.planification.filiere.nom_filiere,
            "module": seance.planification.module,
            "salle": seance.planification.salle.nom_salle,
            "heure_debut": str(seance.planification.heure_debut),
            "heure_fin": str(seance.planification.heure_fin),
        },
        "total_presents": presences.count(),
        "presences": [
            {
                "nom": p.student.last_name,
                "prenom": p.student.first_name,
                "CIE": getattr(p.student, 'student_profile', None).CIE if hasattr(p.student, 'student_profile') else '',
                "heure_scan": str(p.heure_scan),
                "date_scan": str(p.date_scan),
                "ip_address": p.ip_address,
            }
            for p in presences
        ]
    }

    return Response(data, status=200)


import openpyxl
from io import BytesIO

@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def export_seance_excel(request, seance_id):
    try:
        raw_token     = request.headers.get('Authorization', '').split(' ')[1]
        payload       = jwt.decode(raw_token, settings.SECRET_KEY, algorithms=['HS256'])
        user_id = payload.get('user_id')
        roles = payload.get('roles', [])
        if isinstance(roles, str):
            roles = [roles]
    except Exception:
        return Response({'error': 'Token invalide'}, status=401)

    try:
        seance = SeanceQR.objects.select_related(
            'planification__professor',
            'planification__filiere',
            'planification__salle'
        ).get(id=seance_id)
    except SeanceQR.DoesNotExist:
        return Response({'error': 'Seance introuvable'}, status=404)

    if has_role(roles, 'professor') and seance.planification.professor.id == user_id:
        pass
    elif has_role(roles, 'admin'):
        pass
    elif has_role(roles, 'head_of_department'):
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Acces non autorise'}, status=403)
        hod = getattr(user, 'hod_profile', None)
        if not hod or not hod.department_id:
            return Response({'error': 'Acces non autorise'}, status=403)
        if seance.planification.filiere.departement_id != hod.department_id:
            return Response({'error': 'Seance hors de votre departement'}, status=403)
    elif has_role(roles, 'program_manager'):
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Acces non autorise'}, status=403)
        pm = getattr(user, 'program_manager_profile', None)
        if not pm or not pm.program_id:
            return Response({'error': 'Acces non autorise'}, status=403)
        if seance.planification.filiere_id != pm.program_id:
            return Response({'error': 'Seance hors de votre filiere'}, status=403)
    else:
        return Response({'error': 'Acces non autorise'}, status=403)

    filiere        = seance.planification.filiere
    student_profiles = StudentProfile.objects.filter(
        program=filiere
    ).select_related('user').order_by('user__last_name', 'user__first_name')

    presents_ids = set(
        Presence.objects.filter(seance=seance)
        .values_list('student_id', flat=True)
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Seance {seance_id}"

    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    title_font   = Font(bold=True, size=14, color="FFFFFF")
    title_fill   = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    info_font    = Font(bold=True, size=11)
    header_font  = Font(bold=True, size=11, color="FFFFFF")
    header_fill  = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    fill_present = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    fill_absent  = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    thin_border  = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal='center', vertical='center')

    ws.merge_cells('A1:F1')
    c = ws['A1']
    c.value = f"FEUILLE DE PRESENCE - {filiere.nom_filiere}"
    c.font = title_font; c.fill = title_fill; c.alignment = center_align

    info_data = [
        ('Filiere',  filiere.nom_filiere),
        ('Module',   seance.planification.module),
        ('Salle',    str(seance.planification.salle)),
        ('Date',     str(seance.date)),
        ('Semestre', filiere.semester_filiere),
    ]
    for i, (label, value) in enumerate(info_data, start=2):
        ws.cell(row=i, column=1, value=label).font = info_font
        ws.cell(row=i, column=2, value=value)

    header_row_num = len(info_data) + 2
    headers = ["N", "Nom", "Prenom", "CIE", "Filiere", "Presence"]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row_num, column=col, value=h)
        c.font = header_font; c.fill = header_fill
        c.alignment = center_align; c.border = thin_border

    for i, sp in enumerate(student_profiles, start=1):
        est_present = sp.user.id in presents_ids
        row_num = header_row_num + i
        values = [i, sp.user.last_name, sp.user.first_name, sp.CIE, str(sp.program), 'Present' if est_present else 'Absent']
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=row_num, column=col, value=v)
            c.border = thin_border
            if col == 6:
                c.fill = fill_present if est_present else fill_absent
                c.alignment = center_align

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 22
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 12

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"presences_seance_{seance_id}_{seance.date}.xlsx"
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    log_audit(request, 'EXPORT', f"Export séance #{seance_id} - {filename}", 'SeanceQR', seance_id)
    return response


@api_view(['POST'])
@authentication_classes([])
@permission_classes([AllowAny])
def scanner_presence(request):
    token_uuid = request.data.get('token')
    cie        = request.data.get('CIE')
    latitude_etudiant = request.data.get('latitude_etudiant')
    longitude_etudiant= request.data.get('longitude_etudiant')

    if not all([token_uuid, cie]):
        return Response({'error': 'token et CIE obligatoires'}, status=400)

    x_forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
    if x_forwarded_for:
        ip = x_forwarded_for.split(',')[0].strip()
    else:
        ip = request.META.get('REMOTE_ADDR')

    try:
        qrcode = QRCode.objects.get(token=token_uuid)
    except QRCode.DoesNotExist:
        return Response({'error': 'QR Code invalide'}, status=404)

    if not qrcode.is_valid():
        return Response({'error': 'QR Code expire'}, status=400)

    seance = qrcode.seance

    if latitude_etudiant and longitude_etudiant and seance.latitude_prof and seance.longitude_prof:
        distance = calculer_distance(
            float(seance.latitude_prof), float(seance.longitude_prof),
            float(latitude_etudiant), float(longitude_etudiant)
        )
        if distance > 150:
            return Response({
                'error': f'Vous etes trop loin du QR Code ({int(distance)}m). Vous devez etre dans la salle.'
            }, status=403)

    if Presence.objects.filter(seance=seance, ip_address=ip).exists():
        return Response({'error': 'Ce telephone a deja ete utilise pour cette seance'}, status=400)

    try:
        student_profile = StudentProfile.objects.get(CIE=cie)
    except StudentProfile.DoesNotExist:
        return Response({'error': 'CIE introuvable'}, status=404)

    if student_profile.program_id and student_profile.program_id != seance.planification.filiere_id:
        return Response({
            'error': 'Vous n\'etes pas inscrit dans cette filiere',
            'filiere_etudiant': student_profile.program.nom_filiere if student_profile.program else '',
            'filiere_seance': seance.planification.filiere.nom_filiere,
        }, status=403)

    if Presence.objects.filter(student=student_profile.user, seance=seance).exists():
        return Response({'error': 'Presence deja enregistree'}, status=400)

    now_utc = timezone.now()
    now_local = timezone.localtime(now_utc)

    presence = Presence.objects.create(
        student=student_profile.user,
        seance=seance,
        date_scan=now_local.date(),
        heure_scan=now_local.time(),
        ip_address=ip,
    )

    log_audit(request, 'RECORD_PRESENCE',
        f"Présence enregistrée - {student_profile.user.first_name} {student_profile.user.last_name} "
        f"(CIE: {cie}) pour séance #{seance.id}",
        'Presence', presence.id)

    return Response({
        'status': 'success',
        'message': f'Presence enregistree pour {student_profile.user.first_name} {student_profile.user.last_name}',
        'heure': str(now_local.time())[:8],
    })


import cloudinary.uploader

@api_view(['POST', 'DELETE'])
@authentication_classes([])
@permission_classes([AllowAny])
def update_profile_photo(request):
    roles, user_id = get_role_from_request(request)
    if not roles or not user_id:
        return Response({'error': 'Token invalide'}, status=401)

    try:
        user = User.objects.get(id=user_id)
    except User.DoesNotExist:
        return Response({'error': 'Utilisateur introuvable'}, status=404)

    if request.method == 'DELETE':
        user.photo_url = None
        user.save()
        return Response({'message': 'Photo supprimee'}, status=200)

    photo = request.FILES.get('photo')
    if not photo:
        return Response({'error': 'Aucune image fournie'}, status=400)

    allowed_types = ['image/jpeg', 'image/png', 'image/webp']
    if photo.content_type not in allowed_types:
        return Response({'error': 'Format non supporte (jpeg, png, webp)'}, status=status.HTTP_400_BAD_REQUEST)

    if photo.size > 2 * 1024 * 1024:
        return Response({'error': 'Image (max 2 MB)'}, status=status.HTTP_400_BAD_REQUEST)

    try:
        role_name = roles[0] if roles else 'user'
        result = cloudinary.uploader.upload(
            photo,
            folder=f'presenceqr/profiles/{role_name}s/',
            public_id=f'{role_name}_{user_id}',
            overwrite=True,
            transformation=[
                {'width': 300, 'height': 300, 'crop': 'fill', 'gravity': 'face'},
            ]
        )
        url = result['secure_url']
        user.photo_url = url
        user.save()
        return Response({'message': 'Photo mise a jour', 'photo_url': url}, status=200)
    except Exception as e:
        return Response({'error': f'Erreur Cloudinary : {str(e)}'}, status=500)


@api_view(['GET', 'PATCH'])
@authentication_classes([])
@permission_classes([AllowAny])
def my_profile(request):
    roles, user_id = get_role_from_request(request)
    if not roles or not user_id:
        return Response({'error': 'Token invalide'}, status=401)

    try:
        user = User.objects.get(id=user_id)
    except User.DoesNotExist:
        return Response({'error': 'Utilisateur introuvable'}, status=404)

    student_profile = getattr(user, 'student_profile', None)
    professor_profile = getattr(user, 'professor_profile', None)

    if has_role(roles, 'student'):
        if request.method == 'PATCH':
            return Response({'error': 'Modification non autorisee'}, status=403)
        return Response({
            'id': user.id,
            'nom': user.last_name,
            'prenom': user.first_name,
            'email': user.email,
            'CIE': student_profile.CIE if student_profile else None,
            'filiere': str(student_profile.program) if student_profile and student_profile.program else None,
            'photo_url': user.photo_url,
            'role': 'etudiant',
        })

    if has_role(roles, 'professor'):
        if request.method == 'PATCH':
            return Response({'error': 'Modification non autorisee'}, status=403)
        return Response({
            'id': user.id,
            'nom': user.last_name,
            'prenom': user.first_name,
            'email': user.email,
            'specialization': professor_profile.specialization if professor_profile else None,
            'photo_url': user.photo_url,
            'role': 'enseignant',
        })

    if has_role(roles, 'admin'):
        if request.method == 'PATCH':
            for field in ['last_name', 'first_name', 'email']:
                if field in request.data:
                    setattr(user, field, request.data[field])
            user.save()
        return Response({
            'id': user.id,
            'nom': user.last_name,
            'prenom': user.first_name,
            'email': user.email,
            'photo_url': user.photo_url,
            'role': 'admin',
        })

    return Response({'error': 'Role non reconnu'}, status=400)


import math

def calculer_distance(lat1, lon1, lat2, lon2):
    R = 6371000
    phi1 = math.radians(lat1)
    phi2 = math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dlambda = math.radians(lon2 - lon1)

    a = math.sin(dphi/2)**2 + math.cos(phi1) * math.cos(phi2) * math.sin(dlambda/2)**2
    return R * 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))


@api_view(['POST'])
@authentication_classes([])
@permission_classes([AllowAny])
def mes_planifications(request):
    roles, professor_id = get_role_from_request(request)
    if not has_role(roles, 'professor'):
        return Response({'error': 'Reserve aux enseignants'}, status=403)

    data = request.data.copy()
    data['professor'] = professor_id

    filiere_id = data.get('filiere')
    semestre_id = data.get('semestre')
    module_ref_id = data.get('module_ref')
    if semestre_id and filiere_id:
        if not Semestre.objects.filter(id=semestre_id, filiere_id=filiere_id).exists():
            return Response({'error': 'Semestre invalide pour cette filiere'}, status=400)
    if module_ref_id and semestre_id:
        if not Module.objects.filter(id=module_ref_id, semestre_id=semestre_id).exists():
            return Response({'error': 'Module invalide pour ce semestre'}, status=400)

    serializer = PlanificationSerializer(data=data)
    if serializer.is_valid():
        serializer.save()
        return Response(serializer.data, status=201)
    return Response(serializer.errors, status=400)


@api_view(['PATCH', 'DELETE'])
@authentication_classes([])
@permission_classes([AllowAny])
def mes_planifications_detail(request, planification_id):
    roles, professor_id = get_role_from_request(request)
    if not has_role(roles, 'professor'):
        return Response({'error': 'Reserve aux enseignants'}, status=403)

    try:
        planification = Planification.objects.get(id=planification_id)
    except Planification.DoesNotExist:
        return Response({'error': 'Planification introuvable'}, status=404)

    if planification.professor.id != professor_id:
        return Response({'error': 'Acces non autorise a cette planification'}, status=403)

    if request.method == 'DELETE':
        planification.delete()
        return Response({'message': 'Planification supprimee'}, status=200)

    if request.method == 'PATCH':
        data = request.data.copy()
        data.pop('professor', None)

        serializer = PlanificationSerializer(planification, data=data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=200)
        return Response(serializer.errors, status=400)


@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def emploi_du_temps_etudiant(request):
    roles, user_id = get_role_from_request(request)
    if not has_role(roles, 'student'):
        return Response({'error': 'Reserve aux etudiants'}, status=403)

    try:
        user = User.objects.get(id=user_id)
    except User.DoesNotExist:
        return Response({'error': 'Etudiant introuvable'}, status=404)

    student_profile = getattr(user, 'student_profile', None)
    if not student_profile or not student_profile.program:
        return Response({'error': 'Aucune filiere associee'}, status=404)

    planifications = Planification.objects.filter(
        filiere=student_profile.program
    ).select_related('professor', 'salle', 'filiere')

    ordre_jours = {
        'lundi': 1, 'mardi': 2, 'mercredi': 3,
        'jeudi': 4, 'vendredi': 5, 'samedi': 6, 'dimanche': 7
    }
    planifications = sorted(
        planifications,
        key=lambda p: (ordre_jours.get(p.jour.lower(), 99), p.heure_debut)
    )

    data = {
        "filiere": str(student_profile.program),
        "total": len(planifications),
        "planifications": [
            {
                "id": p.id,
                "jour": p.jour,
                "heure_debut": str(p.heure_debut),
                "heure_fin": str(p.heure_fin),
                "module": p.module,
                "salle": p.salle.nom_salle,
                "enseignant": f"{p.professor.first_name} {p.professor.last_name}",
            }
            for p in planifications
        ]
    }

    return Response(data, status=200)


@api_view(['POST'])
@authentication_classes([])
@permission_classes([AllowAny])
def importer_etudiants(request):
    """Import students from an Excel file (.xlsx)."""
    excel_file = request.FILES.get('file')
    if not excel_file:
        return Response({'error': 'Aucun fichier fourni'}, status=400)

    if not excel_file.name.endswith(('.xlsx', '.xls')):
        return Response({'error': 'Format de fichier invalide. Utilisez un fichier .xlsx'}, status=400)

    try:
        wb = openpyxl.load_workbook(excel_file, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        return Response({'error': f'Erreur de lecture du fichier: {str(e)}'}, status=400)

    if len(rows) < 2:
        return Response({'error': 'Le fichier est vide ou ne contient que l\'en-tête'}, status=400)

    header = [str(cell).strip().lower() if cell else '' for cell in rows[0]]

    col_map = {}
    expected_cols = ['nom', 'prenom', 'email', 'password', 'cie', 'filiere_id']
    for col in expected_cols:
        try:
            col_map[col] = header.index(col)
        except ValueError:
            return Response({'error': f'Colonne "{col}" manquante dans le fichier'}, status=400)

    try:
        col_map['groupe'] = header.index('groupe')
    except ValueError:
        col_map['groupe'] = None

    try:
        col_map['module_ids'] = header.index('module_ids')
    except ValueError:
        col_map['module_ids'] = None

    success_count = 0
    errors = []
    created_students = []

    for i, row in enumerate(rows[1:], start=2):
        try:
            nom = str(row[col_map['nom']]).strip() if row[col_map['nom']] else ''
            prenom = str(row[col_map['prenom']]).strip() if row[col_map['prenom']] else ''
            email = str(row[col_map['email']]).strip().lower() if row[col_map['email']] else ''
            password = str(row[col_map['password']]) if row[col_map['password']] else ''
            cie = str(row[col_map['cie']]).strip() if row[col_map['cie']] else ''
            filiere_id = row[col_map['filiere_id']]

            if not nom or not prenom or not email or not password or not cie or filiere_id is None:
                errors.append({'row': i, 'message': 'Champs obligatoires manquants'})
                continue

            if User.objects.filter(email=email).exists():
                errors.append({'row': i, 'message': f'Email déjà utilisé: {email}'})
                continue

            if StudentProfile.objects.filter(CIE=cie).exists():
                errors.append({'row': i, 'message': f'CNE déjà existant: {cie}'})
                continue

            try:
                filiere_id = int(filiere_id)
            except (ValueError, TypeError):
                errors.append({'row': i, 'message': f'filiere_id invalide: {filiere_id}'})
                continue

            module_ids = []
            if col_map['module_ids'] is not None and row[col_map['module_ids']]:
                raw = str(row[col_map['module_ids']]).strip()
                for mid in raw.split(','):
                    mid = mid.strip()
                    if mid:
                        try:
                            module_ids.append(int(mid))
                        except ValueError:
                            errors.append({'row': i, 'message': f'module_ids invalide: {mid}'})
                            continue

            user = User.objects.create_user(
                email=email, password=password,
                first_name=prenom, last_name=nom,
            )
            from users.views import create_role_assignments, create_profile_for_role
            create_role_assignments(user, [Role.STUDENT])
            create_profile_for_role(user, Role.STUDENT, {
                'CIE': cie,
                'program_id': filiere_id,
            })

            if module_ids:
                profile = getattr(user, 'student_profile', None)
                if profile:
                    profile.modules.set(module_ids)

            created_students.append(user)
            success_count += 1

        except Exception as e:
            errors.append({'row': i, 'message': str(e)})

    log_audit(request, 'IMPORT_STUDENTS',
        f"Importation étudiants: {success_count}/{len(rows) - 1} réussis, {len(errors)} erreurs",
        None, None)

    return Response({
        'success': success_count,
        'total': len(rows) - 1,
        'errors': errors,
        'students_created': StudentUserSerializer(created_students, many=True).data,
    })


@api_view(['POST'])
@authentication_classes([])
@permission_classes([AllowAny])
def importer_enseignants(request):
    """Import teachers from an Excel file (.xlsx)."""
    excel_file = request.FILES.get('file')
    if not excel_file:
        return Response({'error': 'Aucun fichier fourni'}, status=400)

    if not excel_file.name.endswith(('.xlsx', '.xls')):
        return Response({'error': 'Format de fichier invalide. Utilisez un fichier .xlsx'}, status=400)

    try:
        wb = openpyxl.load_workbook(excel_file, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        return Response({'error': f'Erreur de lecture du fichier: {str(e)}'}, status=400)

    if len(rows) < 2:
        return Response({'error': 'Le fichier est vide ou ne contient que l\'en-tête'}, status=400)

    header = [str(cell).strip().lower() if cell else '' for cell in rows[0]]

    col_map = {}
    expected_cols = ['nom', 'prenom', 'email', 'password']
    for col in expected_cols:
        try:
            col_map[col] = header.index(col)
        except ValueError:
            return Response({'error': f'Colonne "{col}" manquante dans le fichier'}, status=400)

    try:
        col_map['specialization'] = header.index('specialization')
    except ValueError:
        col_map['specialization'] = None

    try:
        col_map['department_ids'] = header.index('department_ids')
    except ValueError:
        col_map['department_ids'] = None

    success_count = 0
    errors = []
    created_teachers = []

    for i, row in enumerate(rows[1:], start=2):
        try:
            nom = str(row[col_map['nom']]).strip() if row[col_map['nom']] else ''
            prenom = str(row[col_map['prenom']]).strip() if row[col_map['prenom']] else ''
            email = str(row[col_map['email']]).strip().lower() if row[col_map['email']] else ''
            password = str(row[col_map['password']]) if row[col_map['password']] else ''

            if not nom or not prenom or not email or not password:
                errors.append({'row': i, 'message': 'Champs obligatoires manquants'})
                continue

            if User.objects.filter(email=email).exists():
                errors.append({'row': i, 'message': f'Email déjà utilisé: {email}'})
                continue

            specialization = ''
            if col_map['specialization'] is not None and row[col_map['specialization']]:
                specialization = str(row[col_map['specialization']]).strip()

            department_ids = []
            if col_map['department_ids'] is not None and row[col_map['department_ids']]:
                raw = str(row[col_map['department_ids']]).strip()
                for did in raw.split(','):
                    did = did.strip()
                    if did:
                        try:
                            department_ids.append(int(did))
                        except ValueError:
                            errors.append({'row': i, 'message': f'department_ids invalide: {did}'})
                            continue

            user = User.objects.create_user(
                email=email, password=password,
                first_name=prenom, last_name=nom,
            )
            from users.views import create_role_assignments, create_profile_for_role
            extra_data = {'specialization': specialization}
            if department_ids:
                extra_data['department_ids'] = department_ids
            create_role_assignments(user, [Role.PROFESSOR])
            create_profile_for_role(user, Role.PROFESSOR, extra_data)

            created_teachers.append(user)
            success_count += 1

        except Exception as e:
            errors.append({'row': i, 'message': str(e)})

    log_audit(request, 'IMPORT_TEACHERS',
        f"Importation enseignants: {success_count}/{len(rows) - 1} réussis, {len(errors)} erreurs",
        None, None)

    return Response({
        'success': success_count,
        'total': len(rows) - 1,
        'errors': errors,
        'teachers_created': ProfessorUserSerializer(created_teachers, many=True).data,
    })


@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def etudiants_par_module(request):
    """List students grouped by module (via M2M on StudentProfile)."""
    roles, user_id = get_role_from_request(request)
    if not roles or not user_id:
        return Response([], status=200)

    modules = Module.objects.select_related('semestre__filiere').all()
    if has_role(roles, 'head_of_department'):
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response([], status=200)
        hod = getattr(user, 'hod_profile', None)
        if hod and hod.department_id:
            modules = modules.filter(semestre__filiere__departement_id=hod.department_id)
        else:
            return Response([], status=200)
    elif has_role(roles, 'program_manager'):
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response([], status=200)
        pm = getattr(user, 'program_manager_profile', None)
        if pm and pm.program_id:
            modules = modules.filter(semestre__filiere_id=pm.program_id)
        else:
            return Response([], status=200)
    elif not has_role(roles, 'admin'):
        return Response([], status=200)

    result = []
    for module in modules:
        students = User.objects.filter(
            role_assignments__role__codename='student',
            role_assignments__is_active=True,
            student_profile__modules=module,
        ).distinct().order_by('last_name', 'first_name')

        filiere_name = module.semestre.filiere.nom_filiere if module.semestre and module.semestre.filiere else ''

        result.append({
            'module_id': module.id,
            'module': module.nom_module,
            'semestre': module.semestre.nom_semestre,
            'filiere': filiere_name,
            'etudiants': StudentUserSerializer(students, many=True).data,
            'total': students.count(),
        })

    return Response(result)


@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def export_etudiants_par_module_excel(request):
    try:
        roles, user_id = get_role_from_request(request)
        modules = Module.objects.select_related('semestre__filiere').all()
        if has_role(roles, 'head_of_department'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                modules = modules.none()
            hod = getattr(user, 'hod_profile', None)
            if hod and hod.department_id:
                modules = modules.filter(semestre__filiere__departement_id=hod.department_id)
            else:
                modules = modules.none()
        elif has_role(roles, 'program_manager'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                modules = modules.none()
            pm = getattr(user, 'program_manager_profile', None)
            if pm and pm.program_id:
                modules = modules.filter(semestre__filiere_id=pm.program_id)
            else:
                modules = modules.none()
        elif not has_role(roles, 'admin'):
            modules = modules.none()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Etudiants par module"

        from openpyxl.styles import Font, PatternFill, Alignment

        title_font  = Font(bold=True, size=14, color="FFFFFF")
        title_fill  = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        center      = Alignment(horizontal='center', vertical='center')

        ws.merge_cells('A1:H1')
        c = ws['A1']
        c.value = "LISTE DES ETUDIANTS PAR MODULE"
        c.font = title_font; c.fill = title_fill; c.alignment = center

        headers = ["Filiere", "Semestre", "Module", "ID", "Prenom", "Nom", "Email", "CIE"]
        header_row = 3
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col, value=h)
            cell.font = header_font; cell.fill = header_fill
            cell.alignment = center

        row = header_row + 1
        for module in modules:
            students = User.objects.filter(
                role_assignments__role__codename='student',
                role_assignments__is_active=True,
                student_profile__modules=module,
            ).distinct().order_by('last_name', 'first_name')

            filiere_name = module.semestre.filiere.nom_filiere if module.semestre and module.semestre.filiere else ''
            semestre_name = module.semestre.nom_semestre if module.semestre else ''

            for s in students:
                cie = getattr(s.student_profile, 'CIE', '') if hasattr(s, 'student_profile') else ''
                values = [filiere_name, semestre_name, module.nom_module, s.id, s.first_name, s.last_name, s.email, cie]
                for col, v in enumerate(values, start=1):
                    ws.cell(row=row, column=col, value=v)
                row += 1

        for col_letter, w in {'A': 16, 'B': 12, 'C': 24, 'D': 8, 'E': 18, 'F': 18, 'G': 30, 'H': 16}.items():
            ws.column_dimensions[col_letter].width = w

        if row == header_row + 1:
            ws.cell(row=row, column=1, value="Aucun etudiant trouve.")

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        from datetime import date
        filename = f"etudiants_par_module_{date.today().isoformat()}.xlsx"
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        log_audit(request, 'EXPORT', f"Export étudiants par module - {filename}", None, None)
        return response
    except Exception as e:
        return Response({'error': str(e)}, status=500)


@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def absences(request):
    roles, user_id = get_role_from_request(request)
    if not roles or not user_id:
        return Response({'error': 'Authentification requise'}, status=401)

    seances = SeanceQR.objects.select_related(
        'planification__professor',
        'planification__filiere',
        'planification__semestre',
        'planification__module_ref',
        'planification__salle',
    )

    if has_role(roles, 'admin'):
        pass
    elif has_role(roles, 'head_of_department'):
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        hod = getattr(user, 'hod_profile', None)
        if not hod or not hod.department_id:
            return Response({'error': 'Acces non autorise'}, status=403)
        seances = seances.filter(planification__filiere__departement_id=hod.department_id)
    elif has_role(roles, 'program_manager'):
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        pm = getattr(user, 'program_manager_profile', None)
        if not pm or not pm.program_id:
            return Response({'error': 'Acces non autorise'}, status=403)
        seances = seances.filter(planification__filiere_id=pm.program_id)
    elif has_role(roles, 'professor'):
        seances = seances.filter(planification__professor_id=user_id)
    else:
        return Response({'error': 'Acces non autorise'}, status=403)

    filiere_id = request.GET.get('filiere')
    semestre_id = request.GET.get('semestre')
    module_id = request.GET.get('module')
    etudiant_id = request.GET.get('etudiant')
    professeur_id = request.GET.get('professor')
    date_debut = request.GET.get('date_debut')
    date_fin = request.GET.get('date_fin')

    if filiere_id:
        seances = seances.filter(planification__filiere_id=filiere_id)
    if semestre_id:
        seances = seances.filter(planification__semestre_id=semestre_id)
    if module_id:
        seances = seances.filter(planification__module_ref_id=module_id)
    if professeur_id:
        seances = seances.filter(planification__professor_id=professeur_id)
    if date_debut:
        seances = seances.filter(date__gte=date_debut)
    if date_fin:
        seances = seances.filter(date__lte=date_fin)

    seances = seances.order_by('-date', 'planification__heure_debut')

    result = []
    for seance in seances:
        plan = seance.planification
        module_ref = plan.module_ref
        if not module_ref:
            continue

        enrolled = User.objects.filter(
            role_assignments__role__codename='student',
            role_assignments__is_active=True,
            student_profile__modules=module_ref,
        ).select_related('student_profile').order_by('last_name', 'first_name')

        present_qs = Presence.objects.filter(seance=seance)
        present_ids = set(present_qs.values_list('student_id', flat=True))
        present_map = {p.student_id: p for p in present_qs.select_related('student')}

        total_presents = 0
        total_absents = 0
        etudiants_data = []

        for s in enrolled:
            is_present = s.id in present_ids
            if is_present:
                total_presents += 1
            else:
                total_absents += 1

            if etudiant_id and int(etudiant_id) != s.id:
                continue

            p = present_map.get(s.id)
            etudiants_data.append({
                'student_id': s.id,
                'nom': s.last_name,
                'prenom': s.first_name,
                'email': s.email,
                'CIE': getattr(s.student_profile, 'CIE', '') if hasattr(s, 'student_profile') else '',
                'present': is_present,
                'heure_scan': str(p.heure_scan) if p else None,
            })

        if not etudiants_data:
            continue

        result.append({
            'seance_id': seance.id,
            'date': str(seance.date),
            'module': module_ref.nom_module,
            'module_id': module_ref.id,
            'filiere': plan.filiere.nom_filiere,
            'filiere_id': plan.filiere.id,
            'semestre': plan.semestre.nom_semestre if plan.semestre else '',
            'semestre_id': plan.semestre.id if plan.semestre else None,
            'professor': f"{plan.professor.first_name} {plan.professor.last_name}",
            'professor_id': plan.professor.id,
            'heure_debut': str(plan.heure_debut),
            'heure_fin': str(plan.heure_fin),
            'salle': plan.salle.nom_salle,
            'total_presents': total_presents,
            'total_absents': total_absents,
            'total_enrolled': total_presents + total_absents,
            'etudiants': etudiants_data,
        })

    return Response(result)


@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def export_absences_excel(request):
    roles, user_id = get_role_from_request(request)
    if not roles or not user_id:
        return Response({'error': 'Authentification requise'}, status=401)

    seances = SeanceQR.objects.select_related(
        'planification__professor',
        'planification__filiere',
        'planification__semestre',
        'planification__module_ref',
        'planification__salle',
    )

    if has_role(roles, 'admin'):
        pass
    elif has_role(roles, 'head_of_department'):
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        hod = getattr(user, 'hod_profile', None)
        if not hod or not hod.department_id:
            return Response({'error': 'Acces non autorise'}, status=403)
        seances = seances.filter(planification__filiere__departement_id=hod.department_id)
    elif has_role(roles, 'program_manager'):
        try:
            user = User.objects.get(id=user_id)
        except User.DoesNotExist:
            return Response({'error': 'Utilisateur introuvable'}, status=404)
        pm = getattr(user, 'program_manager_profile', None)
        if not pm or not pm.program_id:
            return Response({'error': 'Acces non autorise'}, status=403)
        seances = seances.filter(planification__filiere_id=pm.program_id)
    elif has_role(roles, 'professor'):
        seances = seances.filter(planification__professor_id=user_id)
    else:
        return Response({'error': 'Acces non autorise'}, status=403)

    filiere_id = request.GET.get('filiere')
    semestre_id = request.GET.get('semestre')
    module_id = request.GET.get('module')
    professeur_id = request.GET.get('professor')
    date_debut = request.GET.get('date_debut')
    date_fin = request.GET.get('date_fin')

    if filiere_id:
        seances = seances.filter(planification__filiere_id=filiere_id)
    if semestre_id:
        seances = seances.filter(planification__semestre_id=semestre_id)
    if module_id:
        seances = seances.filter(planification__module_ref_id=module_id)
    if professeur_id:
        seances = seances.filter(planification__professor_id=professeur_id)
    if date_debut:
        seances = seances.filter(date__gte=date_debut)
    if date_fin:
        seances = seances.filter(date__lte=date_fin)

    seances = seances.order_by('-date', 'planification__heure_debut')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rapport des présences"

    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    title_font   = Font(bold=True, size=14, color="FFFFFF")
    title_fill   = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    sub_font     = Font(bold=True, size=11)
    header_font  = Font(bold=True, size=11, color="FFFFFF")
    header_fill  = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    fill_present = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    fill_absent  = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    fill_seance  = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    thin_border  = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal='center', vertical='center')

    ws.merge_cells('A1:P1')
    c = ws['A1']
    c.value = "RAPPORT DES PRESENCES"
    c.font = title_font; c.fill = title_fill; c.alignment = center_align

    headers = [
        "Date", "Jour", "Horaire debut", "Horaire fin", "Module", "Salle",
        "Filiere", "Semestre", "Professeur", "ID etudiant", "Prenom",
        "Nom", "Email", "CIE", "Present", "Heure scan"
    ]

    row = 3
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=16)
    ws.cell(row=row, column=1, value="Filtres appliques:").font = sub_font
    row += 1
    filters_text = []
    if filiere_id: filters_text.append(f"Filiere: {filiere_id}")
    if semestre_id: filters_text.append(f"Semestre: {semestre_id}")
    if module_id: filters_text.append(f"Module: {module_id}")
    if professeur_id: filters_text.append(f"Professeur: {professeur_id}")
    if date_debut: filters_text.append(f"Du: {date_debut}")
    if date_fin: filters_text.append(f"Au: {date_fin}")
    filters_str = ", ".join(filters_text) if filters_text else "Aucun"
    ws.cell(row=row, column=1, value=filters_str)
    row += 2

    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = header_font; c.fill = header_fill
        c.alignment = center_align; c.border = thin_border
    row += 1

    total_seances = 0
    total_presents = 0
    total_absents = 0

    for seance in seances:
        plan = seance.planification
        module_ref = plan.module_ref
        if not module_ref:
            continue

        enrolled = User.objects.filter(
            role_assignments__role__codename='student',
            role_assignments__is_active=True,
            student_profile__modules=module_ref,
        ).select_related('student_profile').order_by('last_name', 'first_name')

        present_qs = Presence.objects.filter(seance=seance)
        present_ids = set(present_qs.values_list('student_id', flat=True))
        present_map = {p.student_id: p for p in present_qs.select_related('student')}

        date_val = str(seance.date)
        try:
            dt = seance.date
            jours = ['Lun', 'Mar', 'Mer', 'Jeu', 'Ven', 'Sam', 'Dim']
            jour_val = jours[dt.weekday()]
        except Exception:
            jour_val = ''
        h_debut = str(plan.heure_debut)[:5]
        h_fin = str(plan.heure_fin)[:5]
        module_val = module_ref.nom_module
        salle_val = plan.salle.nom_salle
        filiere_val = plan.filiere.nom_filiere
        semestre_val = plan.semestre.nom_semestre if plan.semestre else ''
        professeur_val = f"{plan.professor.first_name} {plan.professor.last_name}"

        for s in enrolled:
            is_present = s.id in present_ids
            if is_present:
                total_presents += 1
            else:
                total_absents += 1

            p = present_map.get(s.id)
            heure_scan = str(p.heure_scan)[:5] if p and p.heure_scan else ''

            values = [
                date_val, jour_val, h_debut, h_fin, module_val, salle_val,
                filiere_val, semestre_val, professeur_val, s.id,
                s.first_name, s.last_name, s.email,
                getattr(s.student_profile, 'CIE', '') if hasattr(s, 'student_profile') else '',
                'Oui' if is_present else 'Non', heure_scan,
            ]
            for col, v in enumerate(values, start=1):
                c = ws.cell(row=row, column=col, value=v)
                c.border = thin_border
                if col == 15:
                    c.fill = fill_present if is_present else fill_absent
                    c.alignment = center_align
            row += 1

        total_seances += 1

    widths = [12, 6, 12, 12, 22, 10, 14, 10, 22, 12, 16, 16, 28, 14, 10, 12]
    for i, w in enumerate(widths, start=1):
        letter = chr(64 + i) if i <= 26 else chr(64 + (i - 1) // 26) + chr(65 + (i - 1) % 26)
        ws.column_dimensions[letter].width = w

    if total_seances == 0:
        ws.cell(row=row, column=1, value="Aucune seance trouvee avec les filtres actuels.")

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    from datetime import date
    filename = f"rapport_presences_{date.today().isoformat()}.xlsx"
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    log_audit(request, 'EXPORT', f"Export rapport présences/absences - {filename}", None, None)
    return response


@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def export_etudiants_excel(request):
    try:
        roles, user_id = get_role_from_request(request)
        qs = User.objects.filter(
            role_assignments__role__codename='student',
            role_assignments__is_active=True
        )
        if has_role(roles, 'admin'):
            pass
        elif has_role(roles, 'head_of_department'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                qs = qs.none()
            hod = getattr(user, 'hod_profile', None)
            if hod and hod.department_id:
                qs = qs.filter(student_profile__program__departement_id=hod.department_id)
            else:
                qs = qs.none()
        elif has_role(roles, 'program_manager'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                qs = qs.none()
            pm = getattr(user, 'program_manager_profile', None)
            if pm and pm.program_id:
                qs = qs.filter(student_profile__program_id=pm.program_id)
            else:
                qs = qs.none()
        elif has_role(roles, 'professor'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                qs = qs.none()
            prof = getattr(user, 'professor_profile', None)
            if prof:
                dept_ids = prof.departements.values_list('id', flat=True)
                qs = qs.filter(student_profile__program__departement_id__in=dept_ids)
            else:
                qs = qs.none()
        else:
            qs = qs.none()

        qs = qs.order_by('last_name', 'first_name')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Etudiants"

        from openpyxl.styles import Font, PatternFill, Alignment

        title_font  = Font(bold=True, size=14, color="FFFFFF")
        title_fill  = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        center      = Alignment(horizontal='center', vertical='center')

        ws.merge_cells('A1:G1')
        c = ws['A1']
        c.value = "LISTE DES ETUDIANTS"
        c.font = title_font; c.fill = title_fill; c.alignment = center

        headers = ["ID", "Prenom", "Nom", "Email", "CIE", "Filiere", "Modules"]
        header_row = 3
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col, value=h)
            cell.font = header_font; cell.fill = header_fill
            cell.alignment = center

        row = header_row + 1
        for s in qs:
            sp = getattr(s, 'student_profile', None)
            cie = sp.CIE if sp and sp.CIE else ''
            filiere = sp.program.nom_filiere if sp and sp.program else ''
            if sp:
                modules = sp.modules.all()
                module_str = ', '.join(m.nom_module for m in modules)
            else:
                module_str = ''
            values = [s.id, s.first_name, s.last_name, s.email, cie, filiere, module_str]
            for col, v in enumerate(values, start=1):
                ws.cell(row=row, column=col, value=v)
            row += 1

        for col_letter, w in {'A': 8, 'B': 18, 'C': 18, 'D': 30, 'E': 16, 'F': 20, 'G': 40}.items():
            ws.column_dimensions[col_letter].width = w

        if row == header_row + 1:
            ws.cell(row=row, column=1, value="Aucun etudiant trouve.")

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        from datetime import date
        filename = f"etudiants_{date.today().isoformat()}.xlsx"
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        log_audit(request, 'EXPORT', f"Export liste étudiants - {filename}", None, None)
        return response
    except Exception as e:
        return Response({'error': str(e)}, status=500)


@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def export_modules_excel(request):
    try:
        roles, user_id = get_role_from_request(request)
        qs = Module.objects.select_related('semestre__filiere').all()

        if has_role(roles, 'admin'):
            pass
        elif has_role(roles, 'head_of_department'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                qs = qs.none()
            hod = getattr(user, 'hod_profile', None)
            if hod and hod.department_id:
                qs = qs.filter(semestre__filiere__departement_id=hod.department_id)
            else:
                qs = qs.none()
        elif has_role(roles, 'program_manager'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                qs = qs.none()
            pm = getattr(user, 'program_manager_profile', None)
            if pm and pm.program_id:
                qs = qs.filter(semestre__filiere_id=pm.program_id)
            else:
                qs = qs.none()
        elif has_role(roles, 'professor'):
            try:
                user = User.objects.get(id=user_id)
            except User.DoesNotExist:
                qs = qs.none()
            prof = getattr(user, 'professor_profile', None)
            if prof:
                dept_ids = prof.departements.values_list('id', flat=True)
                qs = qs.filter(semestre__filiere__departement_id__in=dept_ids)
            else:
                qs = qs.none()
        else:
            qs = qs.none()

        qs = qs.order_by('nom_module')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Modules"

        from openpyxl.styles import Font, PatternFill, Alignment

        title_font  = Font(bold=True, size=14, color="FFFFFF")
        title_fill  = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        center      = Alignment(horizontal='center', vertical='center')

        ws.merge_cells('A1:D1')
        c = ws['A1']
        c.value = "LISTE DES MODULES"
        c.font = title_font; c.fill = title_fill; c.alignment = center

        headers = ["ID", "Nom du module", "Semestre", "Filiere"]
        header_row = 3
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col, value=h)
            cell.font = header_font; cell.fill = header_fill
            cell.alignment = center

        row = header_row + 1
        for m in qs:
            filiere_name = m.semestre.filiere.nom_filiere if m.semestre and m.semestre.filiere else ''
            semestre_name = m.semestre.nom_semestre if m.semestre else ''
            values = [m.id, m.nom_module, semestre_name, filiere_name]
            for col, v in enumerate(values, start=1):
                ws.cell(row=row, column=col, value=v)
            row += 1

        for col_letter, w in {'A': 8, 'B': 30, 'C': 16, 'D': 20}.items():
            ws.column_dimensions[col_letter].width = w

        if row == header_row + 1:
            ws.cell(row=row, column=1, value="Aucun module trouve.")

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        from datetime import date
        filename = f"modules_{date.today().isoformat()}.xlsx"
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        log_audit(request, 'EXPORT', f"Export liste modules - {filename}", None, None)
        return response
    except Exception as e:
        return Response({'error': str(e)}, status=500)


@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def debug_whoami(request):
    user = getattr(request, 'user', None)
    if user and user.is_authenticated:
        return Response({
            'email': user.email,
            'id': user.id,
            'roles': user.get_active_roles(),
            'auth_method': 'JWTAuthentication',
        })
    roles, user_id = get_role_from_request(request)
    if user_id:
        try:
            u = User.objects.get(id=user_id)
            return Response({
                'email': u.email,
                'id': u.id,
                'roles': u.get_active_roles(),
                'auth_method': 'manual_jwt_decode',
                'token_roles': roles,
            })
        except User.DoesNotExist:
            return Response({'error': 'user_not_found', 'user_id': user_id, 'roles': roles})
    return Response({
        'error': 'no_auth',
        'auth_header_present': bool(request.headers.get('Authorization', '')),
        'auth_header_prefix': request.headers.get('Authorization', '')[:10] if request.headers.get('Authorization') else None,
    }, status=401)


@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def audit_logs(request):
    roles, user_id = get_role_from_request(request)
    if not has_role(roles, 'admin'):
        return Response({'error': 'Acces reserve aux administrateurs'}, status=403)

    qs = AuditLog.objects.select_related('user').all()

    action_type = request.GET.get('action_type')
    if action_type:
        qs = qs.filter(action_type=action_type)

    user_id_filter = request.GET.get('user_id')
    if user_id_filter:
        qs = qs.filter(user_id=user_id_filter)

    search = request.GET.get('search')
    if search:
        qs = qs.filter(
            Q(description__icontains=search) |
            Q(user__email__icontains=search) |
            Q(user__first_name__icontains=search) |
            Q(user__last_name__icontains=search)
        )

    date_from = request.GET.get('date_from')
    if date_from:
        qs = qs.filter(timestamp__gte=date_from)

    date_to = request.GET.get('date_to')
    if date_to:
        qs = qs.filter(timestamp__lte=date_to + 'T23:59:59')

    page = int(request.GET.get('page', 1))
    page_size = int(request.GET.get('page_size', 50))
    total = qs.count()
    start = (page - 1) * page_size
    end = start + page_size
    qs = qs[start:end]

    serializer = AuditLogSerializer(qs, many=True)
    return Response({
        'total': total,
        'page': page,
        'page_size': page_size,
        'results': serializer.data,
    })


@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def home(request):
    return HttpResponse("Render works!")
